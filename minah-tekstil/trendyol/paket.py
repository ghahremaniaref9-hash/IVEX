"""Trendyol şablonunu XML düzeyinde yazar. Yalnızca ürün satırları değişir; indirilen dosyanın geri kalanı
olduğu gibi kalır: stiller, açılır listeler, özellik sayfaları, yardım sayfası.

openpyxl ile yeniden kaydedilip yardımcı sayfaları çıkarılan dosyayı Trendyol "Excel dosyasında kayıt edilecek
satır bulunamadı" diye reddetti; şablonun bütün parçaları korununca dosya kabul edildi. Metinler paylaşılan metin
tablosuna (sharedStrings.xml) yazılır. Şablonlar farklı araçlarla üretildiği için (Excel, Trendyol'un kendi
üreticisi) XML öznitelikleri sıradan bağımsız okunur.
"""
import re
import zipfile
from html import unescape
from xml.sax.saxutils import escape

from openpyxl.utils import get_column_letter

BASE_OPTIONAL = "37"  # header style of an optional base column (Stok Kodu) in the first template


class Package:
    def __init__(self, path):
        with zipfile.ZipFile(path) as z:
            self.names = z.namelist()
            self.parts = {n: z.read(n) for n in self.names}
        sst = self.text("xl/sharedStrings.xml")
        self.sst_head = sst[:sst.index("<si>")]
        self.si = re.findall(r"<si>.*?</si>", sst, re.S)
        self.index = {}
        for i, si in enumerate(self.si):
            m = re.fullmatch(r'<si><t(?: xml:space="preserve")?>(.*?)</t></si>', si, re.S)
            if m:
                self.index.setdefault(unescape(m.group(1)), i)

    def text(self, name):
        return self.parts[name].decode("utf-8")

    def put(self, name, text):
        if name not in self.parts:
            self.names.append(name)
        self.parts[name] = text.encode("utf-8")

    def drop(self, name):
        self.parts.pop(name, None)
        if name in self.names:
            self.names.remove(name)

    # --- workbook structure -------------------------------------------------------------------------
    def sheets(self):
        """[(name, rId, part)] in tab order."""
        rels = {a["Id"]: a["Target"] for a in tags(self.text("xl/_rels/workbook.xml.rels"), "Relationship")
                if a["Type"].endswith("/worksheet")}
        return [(unescape(a["name"]), a["r:id"], target_part(rels[a["r:id"]]))
                for a in tags(self.text("xl/workbook.xml"), "sheet")]

    def part(self, name):
        return next(p for n, _, p in self.sheets() if n == name)

    def remove_sheet(self, name):
        _, rid, part = next(s for s in self.sheets() if s[0] == name)
        for xml, tag, key, value in (("xl/workbook.xml", "sheet", "r:id", rid),
                                     ("xl/_rels/workbook.xml.rels", "Relationship", "Id", rid),
                                     ("[Content_Types].xml", "Override", "PartName", "/" + part)):
            text = self.text(xml)
            for raw in re.findall(rf"<{tag} [^>]*/>", text):
                if attrs(raw).get(key) == value:
                    text = text.replace(raw, "")
            self.put(xml, text)
        self.drop(part)
        self.drop(part.replace("worksheets/", "worksheets/_rels/") + ".rels")

    # --- cells --------------------------------------------------------------------------------------
    def string(self, text):
        if text not in self.index:
            space = ' xml:space="preserve"' if text != text.strip() or "\n" in text else ""
            self.si.append(f"<si><t{space}>{escape(text)}</t></si>")
            self.index[text] = len(self.si) - 1
        return self.index[text]

    def row(self, r, values, styles):
        cells = []
        for j, v in enumerate(values, start=1):
            col = get_column_letter(j)
            s = f' s="{styles[col]}"' if styles.get(col) else ""
            if v is None or v == "":
                continue
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                cells.append(f'<c r="{col}{r}"{s}><v>{v}</v></c>')
            else:
                cells.append(f'<c r="{col}{r}"{s} t="s"><v>{self.string(str(v))}</v></c>')
        return f'<row r="{r}">{"".join(cells)}</row>'

    def sheet_data(self, header, header_styles, rows, data_styles):
        body = [self.row(1, header, header_styles)]
        body += [self.row(i, values, data_styles) for i, values in enumerate(rows, start=2)]
        return "<sheetData>" + "".join(body) + "</sheetData>"

    def fill_sheet(self, name, header, header_styles, rows, data_styles):
        """Replace a template sheet's rows, keeping its validations; drop the old image hyperlinks."""
        part = self.part(name)
        xml = self.text(part)
        xml = re.sub(r"<sheetData>.*?</sheetData>|<sheetData/>",
                     lambda _: self.sheet_data(header, header_styles, rows, data_styles), xml, count=1, flags=re.S)
        xml = re.sub(r'<dimension ref="[^"]*"/>',
                     f'<dimension ref="A1:{get_column_letter(len(header))}{len(rows) + 1}"/>', xml, count=1)
        xml = re.sub(r"<hyperlinks>.*?</hyperlinks>", "", xml, flags=re.S)
        self.put(part, xml)
        rels = part.replace("worksheets/", "worksheets/_rels/") + ".rels"
        if rels in self.parts:
            kept = [r for r in re.findall(r"<Relationship [^>]*/>", self.text(rels)) if "/hyperlink" not in r]
            if kept:
                self.put(rels, re.sub(r"<Relationship [^>]*/>", "", self.text(rels)).replace(
                    "</Relationships>", "".join(kept) + "</Relationships>"))
            else:
                self.drop(rels)

    def save(self, path):
        refs = sum(len(re.findall(r't="s"', self.text(n))) for n in self.names if n.startswith("xl/worksheets/s"))
        head = re.sub(r'count="\d+" uniqueCount="\d+"', f'count="{refs}" uniqueCount="{len(self.si)}"',
                      self.sst_head)
        self.put("xl/sharedStrings.xml", head + "".join(self.si) + "</sst>")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            for n in self.names:
                z.writestr(n, self.parts[n])


def attrs(tag):
    return dict(re.findall(r'([\w:]+)="([^"]*)"', tag))


def tags(xml, name):
    """Attributes of every <name .../> element, in document order."""
    return [attrs(t) for t in re.findall(rf"<{name} [^>]*/>", xml)]


def target_part(target):
    return target.lstrip("/") if target.startswith("/") else "xl/" + target


def cell_styles(pkg, name, r):
    """Column letter -> style id of the cells in row r of a sheet."""
    m = re.search(rf'<row r="{r}"[^>]*>(.*?)</row>', pkg.text(pkg.part(name)), re.S)
    styles = {}
    for c in re.findall(r"<c [^>]*>", m.group(1) if m else ""):
        a = attrs(c)
        if "s" in a:
            styles[re.match(r"[A-Z]+", a["r"]).group(0)] = a["s"]
    return styles


def header_styles(pkg, name):
    return cell_styles(pkg, name, 1)


def row_styles(pkg, name):
    """Style ids of the first product row; every image column takes Görsel 1's style."""
    styles = cell_styles(pkg, name, 2)
    if "P" in styles:  # Görsel 1 ... Görsel 8 are P ... W
        styles.update({get_column_letter(j): styles["P"] for j in range(16, 24)})
    return styles
